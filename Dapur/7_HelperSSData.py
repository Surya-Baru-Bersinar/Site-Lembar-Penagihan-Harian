import openpyxl


def hapus_kolom_spesifik(nama_file, kolom_yang_dihapus):
    print(f"--> Membuka file '{nama_file}'...")
    wb = openpyxl.load_workbook(nama_file)
    ws = wb.active

    kolom_ditemukan = {}
    for col_idx in range(1, ws.max_column + 1):
        nama_header = ws.cell(row=1, column=col_idx).value
        if nama_header in kolom_yang_dihapus:
            kolom_ditemukan[nama_header] = col_idx

    if not kolom_ditemukan:
        print("--> Peringatan: Tidak ditemukan kolom yang cocok dengan target.")
        return

    indeks_terurut = sorted(kolom_ditemukan.values(), reverse=True)

    for idx in indeks_terurut:
        nama_kolom = [k for k, v in kolom_ditemukan.items() if v == idx][0]
        ws.delete_cols(idx, 1)
        print(f"--> Kolom '{nama_kolom}' (kolom ke-{idx}) berhasil dihapus.")

    wb.save(nama_file)
    print(f"--> Selesai! Perubahan telah disimpan ke '{nama_file}'.")


if __name__ == "__main__":
    FILE_EXCEL = "Piutang_clean_temp.xlsx"
    KOLOM_TARGET = ["SS", "SS.1"]
    hapus_kolom_spesifik(FILE_EXCEL, KOLOM_TARGET)